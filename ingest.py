#!/usr/bin/env python3
"""
Bulk-file intake (BIE batch protocol). Parse a buyer file (CSV / Excel / Word / PDF)
into clean unified rows in data/, with a ZERO-TOKEN lightweight triage score,
processed in batches of 25 that HALT for your approval between batches.
No scraping, no LLM, no credits - pure file-data triage. Deep work happens later
when you run run.py on the resulting data/ rows.

Usage:
  python ingest.py "buyers.xlsx"          # 25 rows, show table, halt for approval, repeat
  python ingest.py "buyers.xlsx" --all    # process every batch without halting
  python ingest.py "buyers.pdf" --batch-size 50
"""
import argparse, csv, io, os, sys
import normalize

CATEGORY_HINTS = ["brass", "metal", "wood", "wooden", "cotton", "paper", "handicraft", "handcraft",
                  "craft", "decor", "home", "gift", "ritual", "pottery", "ceramic", "tableware",
                  "kitchen", "culinary", "hardware", "architectural", "collectible", "interior",
                  "furnish", "homeware", "artisan", "regalos", "kunsthandwerk", "deko"]


# ---------------- readers (return rows using the file's own headers) ----------------
def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt"):
        return _read_csv(path), "csv"
    if ext in (".xlsx", ".xlsm", ".xls"):
        return _read_xlsx(path), "excel"
    if ext == ".docx":
        return _read_docx(path), "word"
    if ext == ".pdf":
        return _read_pdf(path), "pdf"
    raise SystemExit(f"Unsupported file type '{ext}'. Use CSV, XLSX, DOCX, or PDF.")


def _read_csv(path):
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc); break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("latin-1", "replace")
    return list(csv.DictReader(io.StringIO(text)))


def _read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("Excel needs openpyxl:  pip install openpyxl")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    out = []
    for r in rows[1:]:
        if r is None or all(c is None for c in r):
            continue
        out.append({headers[i]: ("" if c is None else str(c)) for i, c in enumerate(r) if i < len(headers)})
    return out


def _read_docx(path):
    try:
        import docx
    except ImportError:
        raise SystemExit("Word needs python-docx:  pip install python-docx")
    d = docx.Document(path)
    out = []
    for tbl in d.tables:
        if len(tbl.rows) < 2:
            continue
        headers = [c.text.strip() for c in tbl.rows[0].cells]
        for row in tbl.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            out.append({headers[i]: cells[i] for i in range(min(len(headers), len(cells)))})
    return out


def _read_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        raise SystemExit("PDF needs pdfplumber:  pip install pdfplumber")
    out = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for tbl in (page.extract_tables() or []):
                if len(tbl) < 2:
                    continue
                headers = [(h or "").strip() for h in tbl[0]]
                for row in tbl[1:]:
                    out.append({headers[i]: (row[i] or "").strip()
                                for i in range(min(len(headers), len(row)))})
    return out


# ---------------- mapping + scoring ----------------
def map_headers(raw_rows):
    headers = []
    for r in raw_rows:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    return {h: normalize._HEADER_TO_FIELD.get(normalize._norm_header(h)) for h in headers}


def to_unified(raw_rows, mapping, source):
    out = []
    for r in raw_rows:
        rec = normalize._blank_record(source)
        for k, v in r.items():
            f = mapping.get(k)
            if f and v and not rec[f]:
                rec[f] = str(v).strip()
        if rec["company_name"] or rec["email"]:
            out.append(rec)
    return out


def lightweight_score(rec, icp):
    score, reasons = 0, []
    country = (rec.get("country") or "").lower()
    if country and any(country in t.lower() or t.lower() in country for t in icp["target_markets"]):
        score += 30; reasons.append("target market")
    blob = " ".join([rec.get("company_name", ""), rec.get("contact_title", ""), rec.get("notes", "")]).lower()
    if any(k in blob for k in icp["buyer_type_keywords"]):
        score += 30; reasons.append("buyer-type")
    if any(k in blob for k in CATEGORY_HINTS):
        score += 20; reasons.append("category hint")
    ev = (8 if rec.get("website") else 0) + (8 if rec.get("email") else 0) + (4 if rec.get("contact_name") else 0)
    score += ev
    if ev:
        reasons.append("contactable")
    score = min(score, 100)
    band = "A" if score >= 80 else "B" if score >= 60 else "C" if score >= 40 else "D"
    needs = "" if (rec.get("email") or rec.get("website")) else "needs enrichment"
    return score, band, ", ".join(reasons) or "thin evidence", needs


# ---------------- main ----------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--all", action="store_true", help="process every batch without halting")
    ap.add_argument("--batch-size", type=int, default=25)
    args = ap.parse_args()
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    import json
    icp = (json.load(open(args.config, encoding="utf-8"))["icp"] if os.path.exists(args.config)
           else {"target_markets": [], "buyer_type_keywords": []})

    src = args.file if os.path.isabs(args.file) else os.path.join(os.getcwd(), args.file)
    if not os.path.exists(src):
        # also try inside data/
        alt = os.path.join("data", os.path.basename(args.file))
        src = alt if os.path.exists(alt) else src
    raw_rows, fmt = read_rows(src)
    if not raw_rows:
        raise SystemExit("No rows found in the file.")

    mapping = map_headers(raw_rows)
    print(f"\nParsed {len(raw_rows)} rows from {os.path.basename(src)} ({fmt}). Header mapping:")
    for h, f in mapping.items():
        print(f"   {h!r:30} -> {f or '(ignored)'}")
    unmapped_company = not any(f == "company_name" for f in mapping.values())
    if unmapped_company:
        print("   ! No column mapped to company_name - check the file's headers.")

    recs = to_unified(raw_rows, mapping, os.path.basename(src))
    for r in recs:
        r["score"], r["band"], r["reason"], r["needs_enrichment"] = lightweight_score(r, icp)

    bs = args.batch_size
    total = len(recs)
    n_batches = (total + bs - 1) // bs
    stem = os.path.splitext(os.path.basename(src))[0]
    out_path = os.path.join("data", f"{stem}_ingested.csv")
    cols = normalize.FIELDS + ["score", "band", "reason", "needs_enrichment"]
    os.makedirs("data", exist_ok=True)

    print(f"\n{total} buyers -> {n_batches} batch(es) of {bs}. Output: {out_path}\n")
    with open(out_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for b in range(n_batches):
            chunk = recs[b * bs:(b + 1) * bs]
            print(f"===== BATCH {b+1} of {n_batches}  (rows {b*bs+1}-{b*bs+len(chunk)}) =====")
            print(f"{'band':<5}{'score':<7}{'country':<12}{'flag':<18}company")
            print("-" * 76)
            for r in chunk:
                w.writerow(r)
                print(f"{r['band']:<5}{r['score']:<7}{(r.get('country') or '-')[:10]:<12}"
                      f"{(r['needs_enrichment'] or '-')[:16]:<18}{r['company_name'][:30]}")
            print()
            if b < n_batches - 1 and not args.all:
                try:
                    input(f"Batch {b+1} of {n_batches} written. Press Enter to approve Batch {b+2} (Ctrl-C to stop)... ")
                except (KeyboardInterrupt, EOFError):
                    print("\nStopped. Processed batches are saved.")
                    break

    a = sum(1 for r in recs if r["band"] == "A"); bb = sum(1 for r in recs if r["band"] == "B")
    need = sum(1 for r in recs if r["needs_enrichment"])
    print(f"Done. {total} buyers triaged -> {out_path}")
    print(f"Bands: A={a} B={bb} C/D={total-a-bb} | needs enrichment: {need}")
    print(f"Next: run  python run.py  to discover domains, find emails, score & draft.")


if __name__ == "__main__":
    main()
